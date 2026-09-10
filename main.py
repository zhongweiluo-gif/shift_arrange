import os
import sys
import json
import openpyxl
from openpyxl.styles import Font, PatternFill
import pulp
from google.auth.transport.requests import Request
from google.oauth2.credentials import Credentials
from google_auth_oauthlib.flow import InstalledAppFlow
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload

# ================= ベースパス取得関数 =================
def get_base_path():
    if getattr(sys, 'frozen', False):
        return os.path.dirname(sys.executable)
    else:
        return os.path.dirname(os.path.abspath(__file__))

# ================= 外部設定ファイルの読み込み =================
config_path = os.path.join(get_base_path(), 'config.json')
try:
    with open(config_path, 'r', encoding='utf-8') as f:
        CONFIG = json.load(f)
except FileNotFoundError:
    print(f"❌ エラー: 設定ファイル [{config_path}] が見つかりません。")
    sys.exit(1)

FILES = CONFIG.get('FILE_SETTINGS', {})
SPREADSHEET_ID = FILES.get('SPREADSHEET_ID', '')
DOWNLOAD_TASKS = [
    (FILES.get('TARGET_MONTH_SHEET', ''), FILES.get('TARGET_LOCAL_FILE', 'Book1.xlsx')),
    (FILES.get('PREV_MONTH_SHEET', ''), FILES.get('PREV_LOCAL_FILE', 'BookLM.xlsx'))
]

SCOPES = ['https://www.googleapis.com/auth/drive.readonly']
CREDENTIALS_FILE = os.path.join(get_base_path(), 'credentials.json')
TOKEN_FILE = os.path.join(get_base_path(), 'token_drive_exporter.json')

# ================= 1. Google Drive API 連携 =================
def get_drive_service():
    if not os.path.exists(CREDENTIALS_FILE):
        raise FileNotFoundError(f'❌ [{CREDENTIALS_FILE}] が見つかりません。')
    creds = None
    if os.path.exists(TOKEN_FILE):
        creds = Credentials.from_authorized_user_file(TOKEN_FILE, SCOPES)
    if not creds or not creds.valid:
        if creds and creds.expired and creds.refresh_token:
            creds.refresh(Request())
        else:
            flow = InstalledAppFlow.from_client_secrets_file(CREDENTIALS_FILE, SCOPES)
            creds = flow.run_local_server(port=0)
        with open(TOKEN_FILE, 'w') as token:
            token.write(creds.to_json())
    return build('drive', 'v3', credentials=creds)

def fetch_and_split_sheets():
    service = get_drive_service()
    temp_excel_path = os.path.join(get_base_path(), '_temp_full_export.xlsx')
    mime_type_xlsx = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    request = service.files().export_media(fileId=SPREADSHEET_ID, mimeType=mime_type_xlsx)
    with open(temp_excel_path, 'wb') as f:
        downloader = MediaIoBaseDownload(f, request)
        done = False
        while not done:
            status, done = downloader.next_chunk()
    for target_sheet, output_filename in DOWNLOAD_TASKS:
        if not target_sheet: continue
        out_path = os.path.join(get_base_path(), output_filename)
        wb = openpyxl.load_workbook(temp_excel_path)
        if target_sheet not in wb.sheetnames: continue
        for sheet_name in wb.sheetnames:
            if sheet_name != target_sheet: del wb[sheet_name]
        wb.save(out_path)
    if os.path.exists(temp_excel_path): os.remove(temp_excel_path)

# ================= 2. PuLP 数理最適化ソルバー =================
def solve_schedule_with_pulp(file_path, prev_month_file_path=None):
    file_path = os.path.join(get_base_path(), file_path)
    if prev_month_file_path:
        prev_month_file_path = os.path.join(get_base_path(), prev_month_file_path)

    print("🚀 【PuLP + HiGHS 数理最適化ソルバー (夜勤サポート対応版)】を起動中...\n")
    if not os.path.exists(file_path):
        print(f"❌ 指定されたファイルが見つかりません: [{file_path}]")
        return

    wb = openpyxl.load_workbook(file_path, data_only=True)
    sheet_name = wb.sheetnames[0]
    ws = wb[sheet_name]

    name_col, team_col, day_start_col, header_row = None, None, None, None
    for r in range(1, 10):
        for c in range(1, 20):
            val = ws.cell(row=r, column=c).value
            if val == '名前': name_col = c; header_row = r
            elif val == 'チーム': team_col = c
            elif str(val) == '1' and header_row and r == header_row: day_start_col = c

    num_days = 0
    if day_start_col and header_row:
        while True:
            v = ws.cell(row=header_row, column=day_start_col + num_days).value
            if isinstance(v, (int, float)) or (isinstance(v, str) and v.isdigit()):
                num_days += 1
            else: break
    if num_days == 0: num_days = 31
    print(f"📅 当月の日数を自動判定しました: {num_days} 日間")

    target_rows, employees, teams, prefilled, red_border_cells = [], [], [], {}, set()

    def check_red_border(cell):
        border = cell.border
        if not border: return False
        for side in [border.left, border.right, border.top, border.bottom]:
            if side and side.color and side.color.rgb:
                if str(side.color.rgb).upper() == 'FFFF0000': return True
        return False

    seen_employees = set()
    for r in range(header_row + 2, ws.max_row + 1):
        team = ws.cell(row=r, column=team_col).value
        raw_name = ws.cell(row=r, column=name_col).value
        name = str(raw_name or '').strip().replace(' ', '').replace('\u3000', '')
        if name in seen_employees: break
        if team in ['NOC', 'CSC'] and name:
            seen_employees.add(name)
            target_rows.append(r)
            employees.append(name)
            teams.append(team)
            emp_idx = len(employees) - 1
            for d in range(num_days):
                c = day_start_col + d
                cell = ws.cell(row=r, column=c)
                val_str = str(cell.value).strip() if cell.value is not None else ''
                has_red = check_red_border(cell)
                if has_red: red_border_cells.add((emp_idx, d))
                if has_red: prefilled[(emp_idx, d)] = val_str if val_str else '休'
                elif val_str != '': prefilled[(emp_idx, d)] = val_str

    prev_month_aug31_night, prev_month_aug31_ake, prev_month_last5_work = set(), set(), {}
    def is_work_shift_name(val):
        val = str(val or '').strip()
        return False if val in ['', '休', '年'] else True

    if prev_month_file_path and os.path.exists(prev_month_file_path):
        wb_lm = openpyxl.load_workbook(prev_month_file_path, data_only=True)
        ws_lm = wb_lm[wb_lm.sheetnames[0]]
        name_col_l, day_start_col_l, header_row_l = None, None, None
        for r in range(1, 10):
            for c in range(1, 20):
                val = ws_lm.cell(row=r, column=c).value
                if val == '名前': name_col_l = c; header_row_l = r
                elif str(val) == '1' and header_row_l and r == header_row_l: day_start_col_l = c
        prev_num_days = 0
        if day_start_col_l and header_row_l:
            while True:
                v = ws_lm.cell(row=header_row_l, column=day_start_col_l + prev_num_days).value
                if isinstance(v, (int, float)) or (isinstance(v, str) and v.isdigit()): prev_num_days += 1
                else: break
        lm_shifts_map = {}
        for r in range(header_row_l + 2, ws_lm.max_row + 1):
            raw_name = ws_lm.cell(row=r, column=name_col_l).value
            name = str(raw_name or '').strip().replace(' ', '').replace('\u3000', '')
            if name and name not in lm_shifts_map:
                shifts_aug = []
                for d_aug in range(prev_num_days):
                    val = ws_lm.cell(row=r, column=day_start_col_l + d_aug).value
                    shifts_aug.append(str(val).strip() if val is not None else '')
                lm_shifts_map[name] = shifts_aug
        for e, name in enumerate(employees):
            aug_shifts = lm_shifts_map.get(name, [''] * prev_num_days)
            aug_last_val = aug_shifts[-1] if len(aug_shifts) > 0 else ''
            if aug_last_val in ['N', 'N設置作業']: prev_month_aug31_night.add(e)
            elif aug_last_val == '明': prev_month_aug31_ake.add(e)
            last5 = aug_shifts[-5:] if len(aug_shifts) >= 5 else [''] * 5
            prev_month_last5_work[e] = [1 if is_work_shift_name(s) else 0 for s in last5]

    def is_red_or_pink_color(cell):
        fill = cell.fill
        if not fill or not fill.fill_type or fill.fill_type == 'none': return False
        color_str = None
        if fill.start_color and fill.start_color.rgb: color_str = str(fill.start_color.rgb).upper()
        elif fill.fgColor and fill.fgColor.rgb: color_str = str(fill.fgColor.rgb).upper()
        if color_str:
            hex_rgb = color_str[-6:]
            if hex_rgb in ['FFFFFF', '000000', '00FFFFFF']: return False
            try:
                r, g, b = int(hex_rgb[0:2], 16), int(hex_rgb[2:4], 16), int(hex_rgb[4:6], 16)
                if r > 180 and (r - g > 15) and (r - b > 15): return True
                if any(kw in hex_rgb for kw in ['F4CC', 'C0CB', 'D9D9', 'ECEC', 'FFC0', 'FFD', 'FFA', 'FFC']): return True
            except ValueError: pass
        if fill.start_color and fill.start_color.theme is not None:
            if fill.start_color.theme in [5, 6, 7, 8, 9]: return True
        return False

    def is_holiday_check(d):
        c = day_start_col + d
        for r_check in range(1, header_row + 3):
            cell = ws.cell(row=r_check, column=c)
            if is_red_or_pink_color(cell): return True
        return False

    num_emp = len(employees)
    holiday_dates = [d for d in range(num_days) if is_holiday_check(d)]
    target_public_rests = len(holiday_dates) if CONFIG.get('USE_DYNAMIC_PUBLIC_REST') else CONFIG.get('FIXED_PUBLIC_REST')

    shifts = [0, 1, 2, 3, 4]  
    trainees = CONFIG.get('TRAINEES', [])
    primary_mentors = CONFIG.get('PRIMARY_MENTORS', [])

    noc_seniors = [e for e in range(num_emp) if teams[e] == 'NOC' and employees[e] not in trainees]
    csc_seniors = [e for e in range(num_emp) if teams[e] == 'CSC' and employees[e] not in trainees]
    all_seniors = noc_seniors + csc_seniors
    trainee_indices = [e for e in range(num_emp) if employees[e] in trainees]
    mentor_indices = [e for e in range(num_emp) if employees[e] in primary_mentors]

    # ✨ プレチェック
    print("🔍 超高速 Python 事前デッドロック検知（プレチェック）を実行中...")
    fatal_errors = []

    for e in range(num_emp):
        rest_count = sum(1 for d in range(num_days) if (e, d) in prefilled and prefilled[(e, d)] in ['休', '年'])
        if rest_count > target_public_rests:
            fatal_errors.append(f"❌ 【公休超過エラー】{employees[e]} (チーム:{teams[e]}): 事前入力された休みが {rest_count} 日あり、当月の公休上限({target_public_rests}日)を超過しています！")

    for d in range(num_days):
        is_hol = d in holiday_dates
        noc_avail = sum(1 for e in noc_seniors if not ((e, d) in prefilled and prefilled[(e, d)] in ['休', '年']))
        csc_avail = sum(1 for e in csc_seniors if not ((e, d) in prefilled and prefilled[(e, d)] in ['休', '年']))
        noc_req = 2 if is_hol else 3 
        csc_req = 1 if is_hol else 2 
        if noc_avail < noc_req:
            fatal_errors.append(f"❌ 【出勤人数不足エラー】第 {d+1} 日: NOCチームの出勤可能人数が {noc_avail} 名しかいません（最低 {noc_req} 名必要です）。")
        if csc_avail < csc_req:
            fatal_errors.append(f"❌ 【出勤人数不足エラー】第 {d+1} 日: CSCチームの出勤可能人数が {csc_avail} 名しかいません（最低 {csc_req} 名必要です）。")

    for e in range(num_emp):
        if e in prev_month_aug31_night and (e, 0) in prefilled and prefilled[(e, 0)] != '明':
            fatal_errors.append(f"❌ 【シフト遷移エラー】{employees[e]}: 前月最終日が夜勤(N)ですが、当月1日が「明」以外で固定されています！")
        if e in prev_month_aug31_ake and (e, 0) in prefilled and prefilled[(e, 0)] not in ['休', '年']:
            fatal_errors.append(f"❌ 【シフト遷移エラー】{employees[e]}: 前月最終日が明け(明)ですが、当月1日が「休/年」以外で固定されています！")
        last5_hist = prev_month_last5_work.get(e, [0]*5)
        for m in range(1, 6):
            if sum(last5_hist[5-(6-m):5]) + sum(1 for di in range(m) if (e, di) in prefilled and is_work_shift_name(prefilled[(e, di)])) > 5:
                fatal_errors.append(f"❌ 【6連勤エラー】{employees[e]}: 前月からの連続出勤により、月初に6連勤が発生する事前入力があります！")
        for d in range(num_days - 1):
            if (e, d) in prefilled and prefilled[(e, d)] in ['N', 'N設置作業']:
                if (e, d+1) in prefilled and prefilled[(e, d+1)] != '明':
                    fatal_errors.append(f"❌ 【シフト遷移エラー】{employees[e]}: 第 {d+1} 日が夜勤(N)ですが、翌日が「明」以外で固定されています！")
            if (e, d) in prefilled and prefilled[(e, d)] == '明':
                if (e, d+1) in prefilled and prefilled[(e, d+1)] not in ['休', '年']:
                    fatal_errors.append(f"❌ 【シフト遷移エラー】{employees[e]}: 第 {d+1} 日が明け(明)ですが、翌日が休み以外で固定されています！")

    if fatal_errors:
        print("\n" + "="*65)
        print("🚨 プレチェック未通過！以下の絶対不可避な矛盾が検出されました（Excelの希望休等を修正してください）:")
        for err in set(fatal_errors):
            print(err)
        print("="*65 + "\n")
        return

    print("✅ プレチェック通過！深刻なデッドロックは見つかりませんでした。モデル構築を開始します...\n")

    def build_model():
        prob = pulp.LpProblem('Shift_Scheduling', pulp.LpMaximize)
        x = pulp.LpVariable.dicts('x', (range(num_emp), range(num_days), shifts), cat=pulp.LpBinary)
        y_d2 = pulp.LpVariable.dicts('y_d2', (mentor_indices, range(num_days)), cat=pulp.LpBinary)
        penalties = []
        objective_terms = []

        shift_map = {'休': 0, 'D': 1, 'D2': 1, 'D3': 1, 'メ': 1, '保全': 1, 'CPN\n設置作業日': 1, 'CPN設置作業日': 1, 'N': 2, 'N設置作業': 2, '明': 3, '年': 4, 'PM年\nor早上がり': 4}

        for e in range(num_emp):
            for d in range(num_days):
                prob += pulp.lpSum([x[e][d][s] for s in shifts]) == 1
                if (e, d) in prefilled:
                    val = prefilled[(e, d)]
                    if val in shift_map:
                        prob += x[e][d][shift_map[val]] == 1
                    else:
                        if (e, d) in red_border_cells: prob += x[e][d][0] == 1 
                        else: prob += x[e][d][1] == 1 
                else:
                    prob += x[e][d][4] == 0

        for e in range(num_emp):
            prob += pulp.lpSum([x[e][d][0] for d in range(num_days)]) == target_public_rests

        for e in range(num_emp):
            if e in prev_month_aug31_night: prob += x[e][0][3] == 1
            elif (e, 0) not in prefilled or prefilled[(e, 0)] != '明': prob += x[e][0][3] == 0
            if e in prev_month_aug31_ake and (e, 0) not in prefilled: prob += x[e][0][0] + x[e][0][4] == 1
            for d in range(num_days - 1):
                prob += x[e][d + 1][3] == x[e][d][2]
                if (e, d + 1) not in prefilled: prob += x[e][d + 1][0] + x[e][d + 1][4] >= x[e][d][3]

        for e in range(num_emp):
            last5_history = prev_month_last5_work.get(e, [0] * 5)
            for m in range(1, 6):
                aug_work_sum = sum(last5_history[5 - (6 - m) : 5])
                prob += aug_work_sum + pulp.lpSum([x[e][d_i][1] + x[e][d_i][2] + x[e][d_i][3] for d_i in range(m)]) <= 5
            for d in range(num_days - 5):
                prob += pulp.lpSum([x[e][d + i][1] + x[e][d + i][2] + x[e][d + i][3] for i in range(6)]) <= 5

        # ✨ [NEW] 每日排班配置 & 夜班支持基础池
        support_needs_names = set(CONFIG.get('NEEDS_NIGHT_SUPPORT', []))
        support_needs_indices = {e for e in range(num_emp) if employees[e] in support_needs_names}
        noc_base_pool = [e for e in noc_seniors if e not in support_needs_indices]
        csc_base_pool = [e for e in csc_seniors if e not in support_needs_indices]

        for d in range(num_days):
            prob += pulp.lpSum([x[e][d][2] for e in noc_base_pool]) == 1
            prob += pulp.lpSum([x[e][d][2] for e in csc_base_pool]) == 1
            prob += pulp.lpSum([x[e][d][2] for e in noc_seniors]) <= 2

        # ✨ [NEW] 夜班搭档(同伴)约束
        support_pool = CONFIG.get('NIGHT_SUPPORT_POOL', {})
        for target_idx in support_needs_indices:
            target_name = employees[target_idx]
            pool_names = support_pool.get(target_name, [])
            valid_pool_indices = [employees.index(name) for name in pool_names if name in employees]
            
            if valid_pool_indices:
                for d in range(num_days):
                    s_support_miss = pulp.LpVariable(f'slack_support_miss_{target_idx}_{d}', cat=pulp.LpBinary)
                    prob += x[target_idx][d][2] - pulp.lpSum([x[p_idx][d][2] for p_idx in valid_pool_indices]) <= s_support_miss
                    penalties.append(CONFIG.get('PENALTY_NIGHT_PARTNER_MISSING', 50000) * s_support_miss)

        # ✨ [NEW] 个人夜班 / 日班 次数限制
        limit_noc = CONFIG.get('LIMIT_NOC_NIGHT_MAX', 4)
        limit_csc = CONFIG.get('LIMIT_CSC_NIGHT_MAX', 6)
        target_support = CONFIG.get('SUPPORT_NIGHT_TARGET', 3)

        for e in noc_seniors:
            prob += pulp.lpSum([x[e][d][1] for d in range(num_days)]) >= 4
            if e in support_needs_indices:
                prob += pulp.lpSum([x[e][d][2] for d in range(num_days)]) == target_support
            else:
                prob += pulp.lpSum([x[e][d][2] for d in range(num_days)]) >= 2
                prob += pulp.lpSum([x[e][d][2] for d in range(num_days)]) <= limit_noc

        for e in csc_seniors:
            if e in support_needs_indices:
                prob += pulp.lpSum([x[e][d][2] for d in range(num_days)]) == target_support
            else:
                prob += pulp.lpSum([x[e][d][2] for d in range(num_days)]) >= 2
                prob += pulp.lpSum([x[e][d][2] for d in range(num_days)]) <= limit_csc

        for e in trainee_indices:
            prob += pulp.lpSum([x[e][d][2] for d in range(num_days)]) >= 2
            prob += pulp.lpSum([x[e][d][2] for d in range(num_days)]) <= 5

        # ✨ [NEW] 夜班平分均衡 (仅在base pool中计算)
        noc_n_max = pulp.LpVariable('noc_n_max', lowBound=0, cat=pulp.LpInteger)
        noc_n_min = pulp.LpVariable('noc_n_min', lowBound=0, cat=pulp.LpInteger)
        for e in noc_base_pool:
            total_n_e = pulp.lpSum([x[e][d][2] for d in range(num_days)])
            prob += total_n_e <= noc_n_max
            prob += total_n_e >= noc_n_min
        penalties.append(CONFIG.get('PENALTY_NIGHT_VARIANCE', 50000) * (noc_n_max - noc_n_min))

        csc_n_max = pulp.LpVariable('csc_n_max', lowBound=0, cat=pulp.LpInteger)
        csc_n_min = pulp.LpVariable('csc_n_min', lowBound=0, cat=pulp.LpInteger)
        for e in csc_base_pool:
            total_n_e = pulp.lpSum([x[e][d][2] for d in range(num_days)])
            prob += total_n_e <= csc_n_max
            prob += total_n_e >= csc_n_min
        penalties.append(CONFIG.get('PENALTY_NIGHT_VARIANCE', 50000) * (csc_n_max - csc_n_min))

        def is_rest_var(e_idx, day_idx):
            if 0 <= day_idx < num_days:
                return x[e_idx][day_idx][0] + x[e_idx][day_idx][4]
            return None

        for e in range(num_emp):
            for d in range(1, num_days - 1):
                rest_prev = is_rest_var(e, d - 1)
                rest_curr = is_rest_var(e, d)
                rest_next = is_rest_var(e, d + 1)
                s_iso_rest = pulp.LpVariable(f'iso_rest_{e}_{d}', cat=pulp.LpBinary)
                prob += rest_curr - rest_prev - rest_next <= s_iso_rest
                penalties.append(CONFIG.get('PENALTY_ISOLATED_REST', 10) * s_iso_rest)
                s_iso_work = pulp.LpVariable(f'iso_work_{e}_{d}', cat=pulp.LpBinary)
                prob += rest_prev + rest_next - rest_curr - 1 <= s_iso_work
                penalties.append(CONFIG.get('PENALTY_ISOLATED_WORK', 10) * s_iso_work)

            for d in range(num_days - 4):
                prefilled_rests = sum(1 for i in range(5) if (e, d+i) in prefilled and prefilled[(e, d+i)] in ['休', '年'])
                if prefilled_rests < 5:
                    s_5_rest = pulp.LpVariable(f'slack_5_rest_{e}_{d}', cat=pulp.LpBinary)
                    prob += pulp.lpSum([is_rest_var(e, d + i) for i in range(5)]) - s_5_rest <= 4
                    penalties.append(CONFIG.get('PENALTY_REST_5_CONSECUTIVE', 100) * s_5_rest)

        special_task_words = {'CPN\n設置作業日', 'CPN設置作業日', '保全', 'メ'}
        for d in range(num_days):
            is_hol = d in holiday_dates
            for m in mentor_indices:
                prob += y_d2[m][d] <= x[m][d][1]
                if (m, d) in prefilled and str(prefilled[(m, d)]).strip() in special_task_words:
                    prob += y_d2[m][d] == 0

            d2_sum = pulp.lpSum([y_d2[m][d] for m in mentor_indices])
            prob += d2_sum == 0 if is_hol else d2_sum == 1

            noc_mentors = [m for m in mentor_indices if m in noc_seniors]
            noc_special_task_count = sum(1 for e in noc_seniors if (e, d) in prefilled and str(prefilled[(e, d)]).strip() in special_task_words)
            noc_pure_d_sum = pulp.lpSum([x[e][d][1] for e in noc_seniors]) - pulp.lpSum([y_d2[m][d] for m in noc_mentors]) - noc_special_task_count
            prob += noc_pure_d_sum == 1 if is_hol else noc_pure_d_sum >= 2
            s_noc_pure_d_over4 = pulp.LpVariable(f'slack_noc_pure_d_over4_{d}', lowBound=0, cat=pulp.LpInteger)
            prob += noc_pure_d_sum - s_noc_pure_d_over4 <= 3
            penalties.append(CONFIG.get('PENALTY_NOC_PURE_D_OVER3', 20000) * s_noc_pure_d_over4)

            if is_hol:
                prob += pulp.lpSum([x[e][d][1] for e in csc_seniors]) == 1
            else:
                prob += pulp.lpSum([x[e][d][1] for e in csc_seniors]) >= 1
                s_csc_d2 = pulp.LpVariable(f'slack_csc_d2_{d}', lowBound=0, cat=pulp.LpInteger)
                prob += pulp.lpSum([x[e][d][1] for e in csc_seniors]) + s_csc_d2 >= 2
                penalties.append(CONFIG.get('PENALTY_CSC_WEEKDAY_D_SHORT', 50000) * s_csc_d2)

            if len(trainee_indices) > 0:
                if is_hol: prob += pulp.lpSum([x[e][d][2] for e in trainee_indices]) == 0
                else: prob += pulp.lpSum([x[e][d][2] for e in trainee_indices]) <= 1

        d2_max = pulp.LpVariable('d2_max', lowBound=0, cat=pulp.LpInteger)
        d2_min = pulp.LpVariable('d2_min', lowBound=0, cat=pulp.LpInteger)
        for m in mentor_indices:
            total_d2_m = pulp.lpSum([y_d2[m][d] for d in range(num_days)])
            prob += total_d2_m <= d2_max
            prob += total_d2_m >= d2_min
        penalties.append(CONFIG.get('PENALTY_D2_RANGE', 20000) * (d2_max - d2_min))

        for m in mentor_indices:
            for d in range(num_days - 2):
                s_d2_dense = pulp.LpVariable(f'slack_d2_dense_{m}_{d}', cat=pulp.LpBinary)
                prob += y_d2[m][d] + y_d2[m][d + 1] + y_d2[m][d + 2] - s_d2_dense <= 1
                penalties.append(CONFIG.get('PENALTY_D2_DENSE', 15000) * s_d2_dense)

        for e in range(num_emp):
            for d in range(num_days):
                for gap in range(1, 6):
                    if d + gap < num_days:
                        s_night_gap = pulp.LpVariable(f'slack_night_gap_{e}_{d}_{gap}', cat=pulp.LpBinary)
                        prob += x[e][d][2] + x[e][d + gap][2] - 1 <= s_night_gap
                        pen_w = CONFIG.get('PENALTY_NIGHT_GAP_TIGHT', 50000) if gap <= 2 else CONFIG.get('PENALTY_NIGHT_GAP_MODERATE', 30000)
                        penalties.append(pen_w * s_night_gap)
                
                start_7 = max(0, d - 6)
                s_7_night = pulp.LpVariable(f'slack_7_night_{e}_{d}', lowBound=0, cat=pulp.LpInteger)
                prob += pulp.lpSum([x[e][d_i][2] for d_i in range(start_7, d + 1)]) - s_7_night <= 2
                penalties.append(CONFIG.get('PENALTY_ROLLING_7DAY_NIGHT_OVER2', 1000) * s_7_night)

        if len(trainee_indices) > 0:
            for d in range(num_days):
                s_d3_over = pulp.LpVariable(f'slack_d3_over_{d}', lowBound=0, cat=pulp.LpInteger)
                prob += pulp.lpSum([x[e][d][1] for e in trainee_indices]) - s_d3_over <= 2
                penalties.append(CONFIG.get('PENALTY_TRAINEE_D3_OVER2', 18000) * s_d3_over)
            for d in holiday_dates:
                for tr in trainee_indices:
                    penalties.append(CONFIG.get('PENALTY_TRAINEE_HOLIDAY_D3', 30000) * x[tr][d][1])

        for e in range(num_emp):
            for d in range(num_days - 1):
                b2 = pulp.LpVariable(f'block_2_{e}_{d}', cat=pulp.LpBinary)
                conds = [is_rest_var(e, d), is_rest_var(e, d + 1)]
                if is_rest_var(e, d - 1) is not None: conds.append(1 - is_rest_var(e, d - 1))
                if is_rest_var(e, d + 2) is not None: conds.append(1 - is_rest_var(e, d + 2))
                for c in conds: prob += b2 <= c
                prob += b2 >= pulp.lpSum(conds) - (len(conds) - 1)
                objective_terms.append(CONFIG.get('REWARD_REST_BLOCK_2', 30) * b2)

            for d in range(num_days - 2):
                b3 = pulp.LpVariable(f'block_3_{e}_{d}', cat=pulp.LpBinary)
                conds = [is_rest_var(e, d), is_rest_var(e, d + 1), is_rest_var(e, d + 2)]
                if is_rest_var(e, d - 1) is not None: conds.append(1 - is_rest_var(e, d - 1))
                if is_rest_var(e, d + 3) is not None: conds.append(1 - is_rest_var(e, d + 3))
                for c in conds: prob += b3 <= c
                prob += b3 >= pulp.lpSum(conds) - (len(conds) - 1)
                objective_terms.append(CONFIG.get('REWARD_REST_BLOCK_3', 40) * b3)

            for d in range(num_days - 3):
                b4 = pulp.LpVariable(f'block_4plus_{e}_{d}', cat=pulp.LpBinary)
                conds = [is_rest_var(e, d), is_rest_var(e, d + 1), is_rest_var(e, d + 2), is_rest_var(e, d + 3)]
                for c in conds: prob += b4 <= c
                prob += b4 >= pulp.lpSum(conds) - (len(conds) - 1)
                objective_terms.append(CONFIG.get('REWARD_REST_BLOCK_4PLUS', 45) * b4)

            for d in holiday_dates:
                objective_terms.append(CONFIG.get('REWARD_HOLIDAY_REST', 60) * (x[e][d][0] + x[e][d][4]))

        prob += pulp.lpSum(objective_terms) - pulp.lpSum(penalties)
        return prob, x, y_d2

    prob, x, y_d2 = build_model()
    threads_count = os.cpu_count() or 4

    print("\n⚡ 高性能 HiGHS ソルバーで計算を実行中 (タイムアウト 10分)...")
    try:
        solver = pulp.HiGHS(msg=True, timeLimit=600, gapRel=0.01, threads=threads_count)
        status = prob.solve(solver)
    except Exception as e:
        print(f"⚠️ HiGHS の呼び出しに失敗しました ({e})。標準 CBC マルチコアモードへフォールバックします...")
        solver = pulp.PULP_CBC_CMD(msg=True, timeLimit=600, gapRel=0.01, threads=threads_count)
        status = prob.solve(solver)

    if pulp.LpStatus[status] in ['Optimal', 'Feasible']:
        print('=' * 65)
        print('🎉 解の導出に成功しました！すべての制約条件を満たす最適シフト表を作成しました。')
        print('=' * 65 + '\n')
    else:
        print('=' * 65)
        print('❌ プレチェックは通過しましたが、複雑な制約の組み合わせにより解が見つかりません。制約を緩めて再度お試しください。')
        print('=' * 65)
        return

    out_file = os.path.join(get_base_path(), FILES.get('PULP_OUTPUT_FILE', 'Book1_PuLP診断版出力結果.xlsx'))
    result_schedule = [['' for _ in range(num_days)] for _ in range(num_emp)]
    val_map = {0: '休', 1: 'D', 2: 'N', 3: '明', 4: '年'}

    for e in range(num_emp):
        for d in range(num_days):
            if (e, d) in prefilled and (e, d) not in red_border_cells:
                result_schedule[e][d] = prefilled[(e, d)]
            else:
                for s in shifts:
                    if pulp.value(x[e][d][s]) is not None and pulp.value(x[e][d][s]) > 0.5:
                        result_schedule[e][d] = val_map[s]

    for d in range(num_days):
        for tr_i in trainee_indices:
            if result_schedule[tr_i][d] == 'D': result_schedule[tr_i][d] = 'D3'
        for m in mentor_indices:
            if pulp.value(y_d2[m][d]) is not None and pulp.value(y_d2[m][d]) > 0.5:
                result_schedule[m][d] = 'D2'

    wb_out = openpyxl.load_workbook(file_path)
    ws_out = wb_out[sheet_name]

    red_font = Font(color='FF0000')
    black_font = Font(color='000000')
    purple_fill = PatternFill(start_color='E6E6FA', end_color='E6E6FA', fill_type='solid')
    green_fill = PatternFill(start_color='90EE90', end_color='90EE90', fill_type='solid')

    for e in range(num_emp):
        r = target_rows[e]
        for d in range(num_days):
            c = day_start_col + d
            val = result_schedule[e][d]
            cell = ws_out.cell(row=r, column=c)
            if (e, d) not in prefilled or (e, d) in red_border_cells:
                cell.value = val
                if val == '休': cell.font = red_font
                elif val in ['N', '明']: cell.fill = purple_fill; cell.font = black_font
                elif val == 'D2': cell.fill = green_fill; cell.font = black_font
                else: cell.font = black_font

    wb_out.save(out_file)
    print(f'📁 シフト表の出力が完了しました: {out_file}\n')


# ================= 3. 役割後処理エンジン (✨ [NEW] D残高保護・純D絶対均等化版) =================
def post_process_schedule(input_file, output_file):
    input_file = os.path.join(get_base_path(), input_file)
    output_file = os.path.join(get_base_path(), output_file)

    print("🚀【シフト表役割後処理エンジン (D残高保護・純D絶対均等化版)】を起動中...\n")

    if not os.path.exists(input_file):
        print(f"❌ 入力ファイルが見つかりません: [{input_file}]")
        return

    wb = openpyxl.load_workbook(input_file)
    sheet_name = wb.sheetnames[0]
    ws = wb[sheet_name]

    name_col, team_col, day_start_col, header_row = None, None, None, None
    for r in range(1, 10):
        for c in range(1, 20):
            val = ws.cell(row=r, column=c).value
            if val == "名前":
                name_col = c; header_row = r
            elif val == "チーム":
                team_col = c
            elif str(val) == "1" and header_row and r == header_row:
                day_start_col = c

    num_days = 0
    if day_start_col and header_row:
        while True:
            v = ws.cell(row=header_row, column=day_start_col + num_days).value
            if isinstance(v, (int, float)) or (isinstance(v, str) and v.isdigit()):
                num_days += 1
            else: break
    if num_days == 0: num_days = 31

    def is_red_or_pink_color(cell):
        fill = cell.fill
        if not fill or not fill.fill_type or fill.fill_type == 'none': return False
        color_str = None
        if fill.start_color and fill.start_color.rgb: color_str = str(fill.start_color.rgb).upper()
        elif fill.fgColor and fill.fgColor.rgb: color_str = str(fill.fgColor.rgb).upper()
        if color_str:
            hex_rgb = color_str[-6:]
            if hex_rgb in ['FFFFFF', '000000', '00FFFFFF']: return False
            try:
                r, g, b = int(hex_rgb[0:2], 16), int(hex_rgb[2:4], 16), int(hex_rgb[4:6], 16)
                if r > 180 and (r - g > 15) and (r - b > 15): return True
                if any(kw in hex_rgb for kw in ['F4CC', 'C0CB', 'D9D9', 'ECEC', 'FFC0', 'FFD', 'FFA', 'FFC']): return True
            except ValueError: pass
        if fill.start_color and fill.start_color.theme is not None:
            if fill.start_color.theme in [5, 6, 7, 8, 9]: return True
        return False

    def is_holiday_check(d):
        c = day_start_col + d
        for r_check in range(1, header_row + 3):
            if is_red_or_pink_color(ws.cell(row=r_check, column=c)): return True
        return False

    holiday_dates = set(d for d in range(num_days) if is_holiday_check(d))

    noc_seniors = []
    csc_employees = []
    seen_employees = set()

    for r in range(header_row + 2, ws.max_row + 1):
        team = ws.cell(row=r, column=team_col).value
        raw_name = ws.cell(row=r, column=name_col).value
        name = str(raw_name or "").strip().replace(" ", "").replace("\u3000", "").replace("\n", "")

        if not name or name in seen_employees: continue
        if team in ["NOC", "CSC"]:
            seen_employees.add(name)
            if team == "NOC":
                noc_seniors.append((name, r))
            elif team == "CSC":
                csc_employees.append((name, r))

    # ================= 2. NOC 初期役割割り当て (残高保護付き基础分配) =================
    noc_matrix = {}
    for name, r in noc_seniors:
        noc_matrix[name] = [str(ws.cell(row=r, column=day_start_col + d).value or "").strip() for d in range(num_days)]

    noc_counts = {name: {"メ": 0, "保全": 0, "メ/保全": 0} for name, _ in noc_seniors}

    for d in range(num_days):
        current_ds = {name: sum(1 for di in range(num_days) if noc_matrix[name][di] == "D") for name, _ in noc_seniors}
        pure_d_candidates = [name for name, _ in noc_seniors if noc_matrix[name][d] == "D"]
        d_count = len(pure_d_candidates)

        if d_count == 2:
            pure_d_candidates.sort(key=lambda x: (noc_counts[x]["メ/保全"], -current_ds[x]))
            noc_matrix[pure_d_candidates[0]][d] = "メ/保全"
            noc_counts[pure_d_candidates[0]]["メ/保全"] += 1

        elif d_count >= 3:
            pure_d_candidates.sort(key=lambda x: (noc_counts[x]["メ"], -current_ds[x]))
            me_name = pure_d_candidates.pop(0)
            noc_matrix[me_name][d] = "メ"
            noc_counts[me_name]["メ"] += 1

            pure_d_candidates.sort(key=lambda x: (noc_counts[x]["保全"], -current_ds[x]))
            ho_name = pure_d_candidates.pop(0)
            noc_matrix[ho_name][d] = "保全"
            noc_counts[ho_name]["保全"] += 1

            # ✨【NEW】D残高が一番「少ない」人をDに残して保護する
            pure_d_candidates.sort(key=lambda x: current_ds[x])
            for name in pure_d_candidates[1:]:
                noc_matrix[name][d] = "タスク"

    # ================= 3. NOC 純D峰値互換平坦化エンジン (削峰填谷: D ⇔ タスク) =================
    print("🔄 [NOC] 純Dの多すぎる人と少なすぎる人の「D ⇔ タスク」直接互換調整を開始します...")
    noc_swap_count = 0

    while True:
        d_totals = {name: sum(1 for d in range(num_days) if noc_matrix[name][d] == "D") for name, _ in noc_seniors}
        if not d_totals: break
        
        max_emp = max(d_totals, key=d_totals.get)
        min_emp = min(d_totals, key=d_totals.get)
        
        if d_totals[max_emp] - d_totals[min_emp] <= 1:
            break

        swapped_in_this_loop = False
        for d in range(num_days):
            if noc_matrix[max_emp][d] == "D":
                task_candidates = [name for name, _ in noc_seniors if noc_matrix[name][d] == "タスク"]
                if task_candidates:
                    task_candidates.sort(key=lambda x: d_totals[x])
                    target_task_emp = task_candidates[0]

                    if d_totals[max_emp] - d_totals[target_task_emp] >= 2:
                        noc_matrix[max_emp][d] = "タスク"
                        noc_matrix[target_task_emp][d] = "D"
                        swapped_in_this_loop = True
                        noc_swap_count += 1
                        print(f"    ⚡ NOC: [{max_emp}](D={d_totals[max_emp]}) ⇔ [{target_task_emp}](D={d_totals[target_task_emp]}) 第 {d+1:2d} 日 (D ⇔ タスク)")
                        break

        if not swapped_in_this_loop:
            print("    ⚠️ NOC: これ以上互換可能な日が存在しないため調整を終了します。")
            break

    print(f"✅ NOC 削峰填谷完了: 合計 {noc_swap_count} 回実行。\n")

    for name, r in noc_seniors:
        for d in range(num_days):
            ws.cell(row=r, column=day_start_col + d).value = noc_matrix[name][d]

    # ================= 4. CSC チーム後処理 (残高保護付き基礎発札 + 再平衡 D ⇔ 勤) =================
    csc_matrix = {name: [str(ws.cell(row=r, column=day_start_col + d).value or "").strip() for d in range(num_days)] for name, r in csc_employees}
    csc_counts = {name: {"メ": 0, "勤": 0} for name, _ in csc_employees}

    for d in range(num_days):
        if d in holiday_dates: continue

        current_ds = {name: sum(1 for di in range(num_days) if csc_matrix[name][di] == "D") for name, _ in csc_employees}
        d_candidates = [name for name, _ in csc_employees if csc_matrix[name][d] == "D"]
        assigned_me = [name for name, _ in csc_employees if csc_matrix[name][d] == "メ"]

        if not assigned_me and len(d_candidates) >= 2:
            d_candidates.sort(key=lambda x: (csc_counts[x]["メ"], -current_ds[x]))
            me_name = d_candidates.pop(0)
            csc_matrix[me_name][d] = "メ"
            csc_counts[me_name]["メ"] += 1

        if d_candidates:
            # ✨【NEW】D残高が一番「少ない」人をDに残して保護する
            d_candidates.sort(key=lambda x: current_ds[x])
            for name in d_candidates[1:]:
                csc_matrix[name][d] = "勤"
                csc_counts[name]["勤"] += 1

    # --- CSC 削峰填谷 (D ⇔ 勤) ---
    print("🔄 [CSC] 純Dの多すぎる人と少なすぎる人の「D ⇔ 勤」直接互換調整を開始します...")
    csc_swap_count = 0

    while True:
        d_totals = {name: sum(1 for d in range(num_days) if csc_matrix[name][d] == "D") for name, _ in csc_employees}
        if not d_totals: break

        max_emp = max(d_totals, key=d_totals.get)
        min_emp = min(d_totals, key=d_totals.get)

        if d_totals[max_emp] - d_totals[min_emp] <= 1:
            break

        swapped_in_this_loop = False
        for d in range(num_days):
            if d in holiday_dates: continue

            if csc_matrix[max_emp][d] == "D":
                kin_candidates = [name for name, _ in csc_employees if csc_matrix[name][d] == "勤"]
                if kin_candidates:
                    kin_candidates.sort(key=lambda x: d_totals[x])
                    target_kin_emp = kin_candidates[0]

                    if d_totals[max_emp] - d_totals[target_kin_emp] >= 2:
                        csc_matrix[max_emp][d] = "勤"
                        csc_matrix[target_kin_emp][d] = "D"
                        swapped_in_this_loop = True
                        csc_swap_count += 1
                        print(f"    ⚡ CSC: [{max_emp}](D={d_totals[max_emp]}) ⇔ [{target_kin_emp}](D={d_totals[target_kin_emp]}) 第 {d+1:2d} 日 (D ⇔ 勤)")
                        break

        if not swapped_in_this_loop:
            print("    ⚠️ CSC: これ以上互換可能な日が存在しないため調整を終了します。")
            break

    print(f"✅ CSC 削峰填谷完了: 合計 {csc_swap_count} 回実行。\n")

    for name, r in csc_employees:
        for d in range(num_days):
            ws.cell(row=r, column=day_start_col + d).value = csc_matrix[name][d]

    # ================= 5. レポートサマリー出力 =================
    noc_report = {name: {"純D": 0, "メ": 0, "保全": 0, "メ/保全": 0, "タスク": 0} for name, _ in noc_seniors}
    for name, _ in noc_seniors:
        for d in range(num_days):
            val = noc_matrix[name][d]
            if val == "D": noc_report[name]["純D"] += 1
            elif val == "メ": noc_report[name]["メ"] += 1
            elif val == "保全": noc_report[name]["保全"] += 1
            elif val == "メ/保全": noc_report[name]["メ/保全"] += 1
            elif val == "タスク": noc_report[name]["タスク"] += 1

    print("📈【NOC 役割分配の最終公平性レポート】:")
    print("-" * 60)
    print(f"{'名前':<10s} | {'純D':<4s} | {'メ':<4s} | {'保全':<4s} | {'メ/保全':<6s} | {'タスク':<6s}")
    print("-" * 60)
    for name, counts in noc_report.items():
        print(f"{name:10s} | {counts['純D']:4d} | {counts['メ']:4d} | {counts['保全']:4d} | {counts['メ/保全']:7d} | {counts['タスク']:4d}")
    print("-" * 60 + "\n")

    csc_report = {name: {"純D": 0, "メ": 0, "勤": 0} for name, _ in csc_employees}
    for name, _ in csc_employees:
        for d in range(num_days):
            val = csc_matrix[name][d]
            if val == "D": csc_report[name]["純D"] += 1
            elif val == "メ": csc_report[name]["メ"] += 1
            elif val == "勤": csc_report[name]["勤"] += 1

    print("📈【CSC 役割分配の最終公平性レポート】:")
    print("-" * 50)
    print(f"{'名前':<10s} | {'純D':<4s} | {'メ':<4s} | {'勤':<4s}")
    print("-" * 50)
    for name, counts in csc_report.items():
        print(f"{name:10s} | {counts['純D']:4d} | {counts['メ']:4d} | {counts['勤']:4d}")
    print("-" * 50 + "\n")

    # 最终样式应用
    red_font = Font(color='FF0000')
    black_font = Font(color='000000')
    purple_fill = PatternFill(start_color='E6E6FA', end_color='E6E6FA', fill_type='solid')
    green_fill = PatternFill(start_color='90EE90', end_color='90EE90', fill_type='solid')

    for r in range(header_row + 2, ws.max_row + 1):
        for d in range(num_days):
            cell = ws.cell(row=r, column=day_start_col + d)
            val = str(cell.value or '').strip()
            if d in holiday_dates and val in ['D2', 'D2(メ)']:
                val = 'D'
                cell.value = 'D'
            if val == '休': cell.font = red_font
            elif val in ['N', '明']: cell.fill = purple_fill; cell.font = black_font
            elif val in ['D2', 'D2(メ)']: cell.fill = green_fill; cell.font = black_font
            else: cell.font = black_font

    wb.save(output_file)
    print("=" * 65)
    print(f"🎉 シフト表の役割削峰填谷処理が完了しました。保存先: [{output_file}]")
    print("=" * 65)


# ================= 4. メインメニュー =================
def main():
    while True:
        print("\n" + "=" * 50)
        print("🌟 スマートシフト作成システム - メインメニュー 🌟")
        print("=" * 50)
        print("  [1] クラウドから最新データを取得")
        print("  [2] PuLP最適化計算を実行 (シフト生成)")
        print("  [3] 役割とルールの後処理 (最终調整・D保護)")
        print("-" * 50)
        print("  [4] 全自動一括実行 (1~3を連続実行)")
        print("  [0] システムを終了")
        print("=" * 50)

        choice = input("👉 実行したいステップの番号を入力してください: ").strip()

        target_file = FILES.get('TARGET_LOCAL_FILE', 'Book1.xlsx')
        prev_file = FILES.get('PREV_LOCAL_FILE', 'BookLM.xlsx')
        pulp_out = FILES.get('PULP_OUTPUT_FILE', 'Book1_PuLP診断版出力結果.xlsx')
        final_out = FILES.get('FINAL_OUTPUT_FILE', 'Book1_最终完成版班表.xlsx')

        if choice == '1':
            try: fetch_and_split_sheets()
            except Exception as e: print(f"❌ エラー: {e}")
        elif choice == '2':
            try: solve_schedule_with_pulp(target_file, prev_file)
            except Exception as e: print(f"❌ エラー: {e}")
        elif choice == '3':
            try: post_process_schedule(pulp_out, final_out)
            except Exception as e: print(f"❌ エラー: {e}")
        elif choice == '4':
            try:
                fetch_and_split_sheets()
                solve_schedule_with_pulp(target_file, prev_file)
                post_process_schedule(pulp_out, final_out)
            except Exception as e: print(f"❌ エラー: {e}")
        elif choice == '0':
            print("👋 プログラムを終了します。")
            break

if __name__ == '__main__':
    main()